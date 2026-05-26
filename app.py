from flask import Flask, request, jsonify, send_file, render_template, redirect
from flask_restx import Api, Resource, fields
import pyodbc
import bcrypt
import jwt
import datetime
import os
from functools import wraps
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from io import BytesIO
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

app = Flask(__name__)
api = Api(app, doc='/docs', title='SIC-NetPOLIx API', version='3.0', 
          description='API completa para el sistema de información de contenido')

# ============================================  
# CONFIGURACIÓN
# ============================================

DB_CONFIG = {
    'DRIVER': '{ODBC Driver 18 for SQL Server}',
    'SERVER': 'localhost',
    'DATABASE': 'SIC_NetPOLIx',
    'Trusted_Connection': 'yes',
    'TrustServerCertificate': 'yes'
}

SECRET_KEY = 'clave_secreta_sic_netpolix_2025!'

def get_db():
    return pyodbc.connect(**DB_CONFIG)

# Decorador para verificar token y rol
def token_required(roles=None):
    def decorator(f):
        @wraps(f)
        def decorated(*args, **kwargs):
            token = request.headers.get('Authorization')
            if not token:
                return {'error': 'Token requerido'}, 401
            try:
                token = token.split(' ')[1]
                payload = jwt.decode(token, SECRET_KEY, algorithms=['HS256'])
                request.user = payload
                
                if roles and payload.get('rol') not in roles:
                    return {'error': 'No tienes permiso'}, 403
                    
            except Exception as e:
                return {'error': f'Token inválido: {str(e)}'}, 401
            return f(*args, **kwargs)
        return decorated
    return decorator

# ============================================
# MODELOS PARA SWAGGER
# ============================================

registro_model = api.model('Registro', {
    'nombre': fields.String(required=True),
    'cedula': fields.String(required=True),
    'email': fields.String(required=True),
    'nickname': fields.String(required=True),
    'password': fields.String(required=True),
    'email_referidor': fields.String(required=False),
    'fecha_nacimiento': fields.String(required=False)
})

login_model = api.model('Login', {
    'email': fields.String(required=True),
    'password': fields.String(required=True)
})


_persona_video_model = api.model('PersonaVideo', {
    'nombre': fields.String(required=True),
    'rol': fields.String(required=True, enum=['ACTOR', 'DIRECTOR', 'PRODUCTOR'])
})

video_model = api.model('Video', {
    'isan': fields.String(required=True),
    'titulo_original': fields.String(required=True),
    'anio': fields.Integer(required=True),
    'duracion': fields.Integer(required=True),
    'descripcion': fields.String,
    'categorias': fields.List(fields.String, description='Máximo 3'),
    'id_clasificacion': fields.String(description='G, PG, PG-13, R o NC-17'),
    'personas': fields.List(fields.Nested(_persona_video_model))
})

calificacion_model = api.model('Calificacion', {
    'valor': fields.String(required=True, enum=['EXCELENTE', 'BUENA', 'REGULAR', 'MALA'])
})

perfil_model = api.model('Perfil', {
    'nombre': fields.String,
    'nickname': fields.String,
    'cedula': fields.String,
    'email': fields.String
})

categoria_model = api.model('Categoria', {
    'nombre': fields.String(required=True)
})

persona_model = api.model('Persona', {
    'nombre': fields.String(required=True),
    'fecha_nacimiento': fields.String(description='Opcional, formato YYYY-MM-DD')
})

# ============================================
# ENDPOINTS PÚBLICOS
# ============================================

@api.route('/api/registro')
class Registro(Resource):
    @api.expect(registro_model)
    def post(self):
        data = request.json
        hashed = bcrypt.hashpw(data['password'].encode('utf-8'), bcrypt.gensalt())
        
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT id_cliente FROM cliente WHERE email = ?", (data['email'],))
            if cursor.fetchone():
                return {'error': 'El email ya está registrado'}, 400
            
            fecha_nac = data.get('fecha_nacimiento', '').strip()
            fecha_nac_hash = bcrypt.hashpw(fecha_nac.encode('utf-8'), bcrypt.gensalt()).decode('utf-8') if fecha_nac else None

            cursor.execute("""
                INSERT INTO cliente (nombre, cedula, email, nickname, password, puntos, saldo, rol, fecha_nacimiento_hash)
                VALUES (?, ?, ?, ?, ?, 0, 0, 'CLIENTE', ?)
            """, (data['nombre'], data['cedula'], data['email'], data['nickname'], hashed.decode('utf-8'), fecha_nac_hash))
            
            cursor.execute("SELECT @@IDENTITY")
            id_cliente = int(cursor.fetchone()[0])
            
            if data.get('email_referidor'):
                cursor.execute("SELECT id_cliente FROM cliente WHERE email = ?", (data['email_referidor'],))
                referidor = cursor.fetchone()
                if referidor:
                    cursor.execute("UPDATE cliente SET puntos = puntos + 1 WHERE id_cliente = ?", (referidor[0],))
                    # También actualizar id_referido en el nuevo cliente
                    cursor.execute("UPDATE cliente SET id_referido = ? WHERE id_cliente = ?", (referidor[0], id_cliente))
            
            conn.commit()
            return {'mensaje': 'Cliente registrado exitosamente', 'id_cliente': id_cliente}, 201
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 400
        finally:
            conn.close()

@api.route('/api/login')
class Login(Resource):
    @api.expect(login_model)
    def post(self):
        data = request.json
        credencial = data.get('email', '').strip()
        password = data.get('password', '')
        conn = get_db()
        cursor = conn.cursor()

        # 1) Buscar en clientes por email
        cursor.execute("""
            SELECT id_cliente, nombre, nickname, cedula, email, password, rol
            FROM cliente WHERE email = ?
        """, (credencial,))
        row = cursor.fetchone()

        if row:
            try:
                pwd_ok = bcrypt.checkpw(password.encode('utf-8'), row[5].encode('utf-8'))
            except ValueError:
                pwd_ok = False
            if not pwd_ok:
                conn.close()
                return {'error': 'Credenciales inválidas'}, 401
            token = jwt.encode({
                'id': row[0], 'nombre': row[1],
                'nickname': row[2] if row[2] else row[1],
                'email': row[4], 'rol': row[6], 'tipo': 'CLIENTE',
                'exp': datetime.datetime.utcnow() + datetime.timedelta(hours=24)
            }, SECRET_KEY, algorithm='HS256')
            conn.close()
            return {
                'token': token, 'nombre': row[1],
                'nickname': row[2] if row[2] else row[1],
                'cedula': row[3], 'email': row[4], 'rol': row[6]
            }

        # 2) Buscar en usuario_sistema por nombre_usuario
        cursor.execute("""
            SELECT id_usuario, nombre_completo, nombre_usuario, documento, password, rol
            FROM usuario_sistema WHERE nombre_usuario = ?
        """, (credencial,))
        row_sys = cursor.fetchone()
        conn.close()

        if not row_sys:
            return {'error': 'Credenciales inválidas - Usuario no encontrado'}, 401

        try:
            pwd_sys_ok = bcrypt.checkpw(password.encode('utf-8'), row_sys[4].encode('utf-8'))
        except ValueError:
            pwd_sys_ok = False
        if not pwd_sys_ok:
            return {'error': 'Credenciales inválidas'}, 401

        token = jwt.encode({
            'id': row_sys[0], 'nombre': row_sys[1],
            'nickname': row_sys[2], 'email': row_sys[2],
            'rol': row_sys[5], 'tipo': 'SISTEMA',
            'exp': datetime.datetime.utcnow() + datetime.timedelta(hours=24)
        }, SECRET_KEY, algorithm='HS256')
        return {
            'token': token, 'nombre': row_sys[1],
            'nickname': row_sys[2], 'rol': row_sys[5]
        }



# ============================================
# ENDPOINTS CLIENTE
# ============================================

@app.route('/pago')
def pago_page():
    return render_template('pago.html')

@api.route('/api/pagar')
class Pagar(Resource):
    @token_required(roles=['CLIENTE'])
    def post(self):
        data = request.json
        items = data.get('items', [])
        metodo_pago = data.get('metodo_pago', 'PUNTOS')  # PUNTOS | TARJETA | SALDO
        if metodo_pago not in ('PUNTOS', 'TARJETA', 'SALDO'):
            return {'error': 'Método de pago inválido. Use PUNTOS, TARJETA o SALDO'}, 400
        id_cliente = request.user.get('id')
        
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT saldo, puntos, fecha_nacimiento_hash FROM cliente WHERE id_cliente = ?", (id_cliente,))
            row = cursor.fetchone()
            if not row:
                return {'error': 'Cliente no encontrado'}, 404

            saldo = float(row[0])
            puntos = int(row[1])
            hash_guardado = row[2]
            total = 0

            fecha_nacimiento = data.get('fecha_nacimiento')
            for item in items:
                es_serie = item.get('es_serie', False)
                item_id = item.get('id') or item.get('id_video')
                if es_serie:
                    cursor.execute("""
                        SELECT DISTINCT v.id_clasificacion
                        FROM serie_video sv
                        JOIN video v ON sv.id_video = v.id_video
                        WHERE sv.id_serie = ? AND v.id_clasificacion IS NOT NULL
                    """, (item_id,))
                    clasificaciones = [r[0] for r in cursor.fetchall()]
                else:
                    cursor.execute("SELECT id_clasificacion FROM video WHERE id_video = ?", (item_id,))
                    r = cursor.fetchone()
                    clasificaciones = [r[0]] if r and r[0] else []
                for clasificacion in clasificaciones:
                    ok, error_edad = verificar_restriccion_edad(clasificacion, fecha_nacimiento, hash_guardado)
                    if not ok:
                        return {'error': error_edad}, 403

            for item in items:
                tipo = item.get('tipo')
                precio = PRECIO_COMPRA if tipo == 'COMPRA' else PRECIO_ALQUILER
                total += precio

            # Validar según método de pago
            if metodo_pago == 'PUNTOS':
                puntos_necesarios = 20 * len(items)
                if puntos < puntos_necesarios:
                    return {'error': f'Puntos insuficientes. Necesitas {puntos_necesarios} puntos y tienes {puntos}'}, 400
            elif metodo_pago == 'SALDO':
                if saldo < total:
                    return {'error': f'Saldo insuficiente. Necesitas ${total:.2f} y tienes ${saldo:.2f}'}, 400

            puntos_ganados = 0
            for item in items:
                tipo_original = item.get('tipo')
                es_serie = item.get('es_serie', False)
                item_id = item.get('id') or item.get('id_video')
                # Al canjear puntos siempre es COMPRA (permanente, sin vencimiento)
                tipo = 'COMPRA' if metodo_pago == 'PUNTOS' else tipo_original
                precio = PRECIO_COMPRA if tipo_original == 'COMPRA' else PRECIO_ALQUILER
                pts = 2 if tipo == 'COMPRA' else 1
                monto = 0 if metodo_pago == 'PUNTOS' else precio

                if es_serie:
                    # Expandir la temporada: crear transacción por cada episodio
                    cursor.execute("SELECT id_video FROM serie_video WHERE id_serie = ?", (item_id,))
                    ids_video = [r[0] for r in cursor.fetchall()]
                else:
                    ids_video = [item_id]

                for id_video in ids_video:
                    if tipo == 'ALQUILER':
                        fecha_exp = (datetime.date.today() + datetime.timedelta(days=30)).isoformat()
                        cursor.execute("""
                            INSERT INTO transaccion (id_cliente, id_video, tipo, monto, fecha_expiracion)
                            VALUES (?, ?, ?, ?, ?)
                        """, (id_cliente, id_video, tipo, monto, fecha_exp))
                    else:
                        cursor.execute("""
                            INSERT INTO transaccion (id_cliente, id_video, tipo, monto)
                            VALUES (?, ?, ?, ?)
                        """, (id_cliente, id_video, tipo, monto))

                if metodo_pago != 'PUNTOS':
                    puntos_ganados += pts

            if metodo_pago == 'PUNTOS':
                cursor.execute("""
                    UPDATE cliente SET puntos = puntos - ?
                    WHERE id_cliente = ?
                """, (20 * len(items), id_cliente))
            elif metodo_pago == 'SALDO':
                cursor.execute("""
                    UPDATE cliente SET saldo = saldo - ?, puntos = puntos + ?
                    WHERE id_cliente = ?
                """, (total, puntos_ganados, id_cliente))
            else:
                cursor.execute("""
                    UPDATE cliente SET puntos = puntos + ?
                    WHERE id_cliente = ?
                """, (puntos_ganados, id_cliente))

            conn.commit()

            cursor.execute("SELECT saldo, puntos FROM cliente WHERE id_cliente = ?", (id_cliente,))
            row = cursor.fetchone()

            return {
                'mensaje': 'Pago procesado exitosamente',
                'total_cobrado': 0 if metodo_pago == 'PUNTOS' else total,
                'puntos_ganados': puntos_ganados,
                'puntos_usados': 20 * len(items) if metodo_pago == 'PUNTOS' else 0,
                'saldo_restante': float(row[0]),
                'puntos_totales': int(row[1])
            }, 200

        except Exception as e:
            conn.rollback()
            print(f"ERROR pagar: {str(e)}")
            return {'error': str(e)}, 500
        finally:
            conn.close()

@api.route('/api/perfil')
class Perfil(Resource):
    @token_required(roles=['CLIENTE'])
    def get(self):
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT id_cliente, nombre, nickname, cedula, email, puntos, saldo, foto_perfil, fecha_nacimiento_hash
            FROM cliente
            WHERE id_cliente = ?
        """, (request.user.get('id'),))
        row = cursor.fetchone()
        conn.close()
        if not row:
            return {'error': 'Cliente no encontrado'}, 404
        return {
            'id': row[0],
            'nombre': row[1],
            'nickname': row[2],
            'cedula': row[3],
            'email': row[4],
            'puntos': row[5],
            'saldo': float(row[6]) if row[6] else 0,
            'foto_perfil': row[7] if row[7] else None,
            'tiene_fecha_nacimiento': row[8] is not None
        }
    
    @token_required(roles=['CLIENTE'])
    @api.expect(perfil_model)
    def put(self):
        data = request.json
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("""
            UPDATE cliente 
            SET nombre = ?, nickname = ?, cedula = ?, email = ?
            WHERE id_cliente = ?
        """, (data.get('nombre'), data.get('nickname'), data.get('cedula'), data.get('email'), request.user.get('id')))
        conn.commit()
        conn.close()
        return {'mensaje': 'Perfil actualizado exitosamente'}
    
@api.route('/api/perfil/foto')
class PerfilFoto(Resource):
    @token_required(roles=['CLIENTE'])
    def put(self):
        data = request.json
        foto = data.get('foto_perfil', '').strip()
        FOTOS_VALIDAS = {f'perfil_{i}.jpg' for i in [1,6,7,9]} | \
                        {f'perfil_{i}.webp' for i in [2,4,5,8]} | \
                        {'perfil_3.jpeg'}
        if foto not in FOTOS_VALIDAS:
            return {'error': 'Foto no válida'}, 400
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("UPDATE cliente SET foto_perfil = ? WHERE id_cliente = ?",
                       (foto, request.user.get('id')))
        conn.commit()
        conn.close()
        return {'mensaje': 'Foto actualizada', 'foto_perfil': foto}

@api.route('/api/perfil/fecha-nacimiento')
class PerfilFechaNacimiento(Resource):
    @token_required(roles=['CLIENTE'])
    def put(self):
        data = request.json
        fecha = data.get('fecha_nacimiento', '').strip()
        if not fecha:
            return {'error': 'Fecha requerida'}, 400
        hashed = bcrypt.hashpw(fecha.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("UPDATE cliente SET fecha_nacimiento_hash = ? WHERE id_cliente = ?",
                       (hashed, request.user.get('id')))
        conn.commit()
        conn.close()
        return {'mensaje': 'Fecha de nacimiento guardada'}

@api.route('/api/cambiar-password')
class CambiarPassword(Resource):
    @token_required(roles=['CLIENTE'])
    def put(self):
        data = request.json
        id_cliente = request.user['id']
        conn = get_db()
        cursor = conn.cursor()
        
        # Obtener contraseña actual
        cursor.execute("SELECT password FROM cliente WHERE id_cliente = ?", (id_cliente,))
        row = cursor.fetchone()
        
        if not row:
            return {'error': 'Usuario no encontrado'}, 404
        
        # Verificar contraseña actual
        if not bcrypt.checkpw(data['current_password'].encode('utf-8'), row[0].encode('utf-8')):
            return {'error': 'Contraseña actual incorrecta'}, 401
        
        # Encriptar nueva contraseña
        hashed = bcrypt.hashpw(data['new_password'].encode('utf-8'), bcrypt.gensalt())
        
        # Actualizar contraseña
        cursor.execute("UPDATE cliente SET password = ? WHERE id_cliente = ?", (hashed.decode('utf-8'), id_cliente))
        conn.commit()
        conn.close()
        
        return {'mensaje': 'Contraseña actualizada exitosamente'}

def get_poster_url(id_video):
    base = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'static', 'posters')
    for ext in ['jpg', 'jpeg', 'png', 'webp']:
        if os.path.exists(os.path.join(base, f'video_{id_video}.{ext}')):
            return f'/static/posters/video_{id_video}.{ext}'
    return None

@api.route('/api/buscar')
class Buscar(Resource):
    @token_required(roles=['CLIENTE'])
    def get(self):
        query = request.args.get('q', '')
        conn = get_db()
        cursor = conn.cursor()
        resultados = {'peliculas': [], 'series': []}

        try:
            # Buscar películas (por título, actor/director o categoría)
            cursor.execute("""
                SELECT DISTINCT v.id_video, v.titulo_original, v.anio, v.duracion,
                       vp.promedio, v.id_clasificacion
                FROM video v
                LEFT JOIN vw_promedio_calificaciones vp ON v.id_video = vp.id_video
                LEFT JOIN video_persona vper ON v.id_video = vper.id_video
                LEFT JOIN persona per ON vper.id_persona = per.id_persona
                LEFT JOIN video_categoria vc ON v.id_video = vc.id_video
                LEFT JOIN categoria c ON vc.id_categoria = c.id_categoria
                WHERE (v.titulo_original LIKE ? OR per.nombre LIKE ? OR c.nombre LIKE ?)
            """, (f'%{query}%', f'%{query}%', f'%{query}%'))

            for row in cursor.fetchall():
                resultados['peliculas'].append({
                    'id': row[0],
                    'titulo': row[1],
                    'anio': row[2],
                    'duracion': row[3],
                    'promedio': float(row[4]) if row[4] else 0,
                    'clasificacion': row[5],
                    'url_imagen': get_poster_url(row[0])
                })

            # Buscar series (por título o categoría de sus episodios)
            cursor.execute("""
                SELECT s.titulo,
                       COUNT(DISTINCT s.id_serie) AS total_temporadas,
                       MIN(s.id_serie) AS id_representativo
                FROM serie s
                LEFT JOIN serie_video sv ON s.id_serie = sv.id_serie
                LEFT JOIN video v ON sv.id_video = v.id_video
                LEFT JOIN video_categoria vc ON v.id_video = vc.id_video
                LEFT JOIN categoria c ON vc.id_categoria = c.id_categoria
                WHERE (s.titulo LIKE ? OR c.nombre LIKE ?)
                GROUP BY s.titulo
            """, (f'%{query}%', f'%{query}%'))

            for row in cursor.fetchall():
                resultados['series'].append({
                    'id': row[2],
                    'titulo': row[0],
                    'total_temporadas': row[1],
                    'url_imagen': get_serie_poster_url(row[2])
                })

            conn.close()
            return resultados

        except Exception as e:
            print(f"Error buscar: {str(e)}")
            conn.close()
            return {'error': str(e), 'peliculas': [], 'series': []}, 500

@api.route('/api/videos')
class ListarVideos(Resource):
    @token_required(roles=['CLIENTE'])
    def get(self):
        categoria = request.args.get('categoria', '')
        titulo = request.args.get('titulo', '')
        conn = get_db()
        cursor = conn.cursor()
        query = """
            SELECT v.id_video, v.titulo_original, v.anio, v.duracion,
                   vp.promedio, v.descripcion,
                   STRING_AGG(per.nombre, ',') WITHIN GROUP (ORDER BY per.nombre) AS directores,
                   STRING_AGG(CASE WHEN vper.rol = 'ACTOR' THEN per.nombre END, ',') WITHIN GROUP (ORDER BY per.nombre) AS actores,
                   STRING_AGG(c.nombre, ',') AS categorias,
                   MAX(v.id_clasificacion) AS clasificacion
            FROM video v
            LEFT JOIN vw_promedio_calificaciones vp ON v.id_video = vp.id_video
            LEFT JOIN video_persona vper ON v.id_video = vper.id_video
            LEFT JOIN persona per ON vper.id_persona = per.id_persona
            LEFT JOIN video_categoria vc ON v.id_video = vc.id_video
            LEFT JOIN categoria c ON vc.id_categoria = c.id_categoria
            WHERE v.tipo = 'PELICULA'
        """
        params = []
        if categoria:
            query += " AND c.nombre = ?"
            params.append(categoria.upper())
        if titulo:
            query += " AND v.titulo_original LIKE ?"
            params.append(f'%{titulo}%')
        query += """
            GROUP BY v.id_video, v.titulo_original, v.anio, v.duracion,
                     vp.promedio, v.descripcion
        """
        cursor.execute(query, params)
        rows = cursor.fetchall()
        conn.close()
        videos = [{
            'id': r[0],
            'titulo': r[1],
            'anio': r[2],
            'duracion': r[3],
            'promedio': float(r[4]) if r[4] else 0,
            'descripcion': r[5],
            'directores': r[6].split(',') if r[6] else [],
            'actores': r[7].split(',') if r[7] else [],
            'categorias': r[8].split(',') if r[8] else [],
            'clasificacion': r[9],
            'url_imagen': get_poster_url(r[0])
        } for r in rows]
        return {'videos': videos}

@api.route('/api/video/<int:id_video>/calificar')
class CalificarVideo(Resource):
    @token_required(roles=['CLIENTE'])
    @api.expect(calificacion_model)
    def post(self, id_video):
        data = request.json
        id_cliente = request.user.get('id')
        valor = data.get('calificacion', '').upper()

        if valor not in ('EXCELENTE', 'BUENA', 'REGULAR', 'MALA'):
            return {'error': 'Valor inválido'}, 400

        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("""
                SELECT id_calificacion FROM calificacion
                WHERE id_video = ? AND id_cliente = ?
            """, (id_video, id_cliente))
            existe = cursor.fetchone()

            if existe:
                cursor.execute("""
                    UPDATE calificacion SET valor = ?, fecha = GETDATE()
                    WHERE id_video = ? AND id_cliente = ?
                """, (valor, id_video, id_cliente))
            else:
                cursor.execute("""
                    INSERT INTO calificacion (id_video, id_cliente, valor)
                    VALUES (?, ?, ?)
                """, (id_video, id_cliente, valor))

            conn.commit()

            cursor2 = conn.cursor()
            cursor2.execute("""
                SELECT AVG(CASE valor
                    WHEN 'EXCELENTE' THEN 4.0
                    WHEN 'BUENA' THEN 3.0
                    WHEN 'REGULAR' THEN 2.0
                    WHEN 'MALA' THEN 1.0
                END)
                FROM calificacion
                WHERE id_video = ?
            """, (id_video,))

            row = cursor2.fetchone()
            promedio = round(float(row[0]), 2) if row and row[0] else 0.0
            return {'mensaje': 'Calificación registrada', 'promedio': promedio}, 200
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 500
        finally:
            conn.close()

@api.route('/api/video/<int:id_video>/calificaciones')
class VerCalificaciones(Resource):
    @token_required(roles=['CLIENTE'])
    def get(self, id_video):
        id_cliente = request.user.get('id')
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT valor, COUNT(*) FROM calificacion WHERE id_video = ? GROUP BY valor", (id_video,))
        rows = cursor.fetchall()
        cursor.execute("SELECT promedio FROM vw_promedio_calificaciones WHERE id_video = ?", (id_video,))
        row_prom = cursor.fetchone()
        cursor.execute("SELECT valor FROM calificacion WHERE id_video = ? AND id_cliente = ?", (id_video, id_cliente))
        mi_voto_row = cursor.fetchone()
        conn.close()
        votos = {"EXCELENTE": 0, "BUENA": 0, "REGULAR": 0, "MALA": 0}
        for row in rows:
            votos[row[0]] = row[1]
        return {
            'promedio': float(row_prom[0]) if row_prom and row_prom[0] else 0,
            'votos': votos,
            'mi_voto': mi_voto_row[0] if mi_voto_row else None
        }

PRECIO_ALQUILER = 3.50
PRECIO_COMPRA = 10.00

EDAD_MINIMA_CLASIFICACION = {'PG-13': 13, 'R': 17, 'NC-17': 18}

def verificar_restriccion_edad(clasificacion, fecha_nacimiento_str, hash_guardado):
    """
    Verifica si el cliente cumple la edad mínima para el contenido.
    No almacena ni registra la fecha de nacimiento — solo la usa en memoria.
    Retorna (True, None) si puede acceder, (False, mensaje) si no.
    """
    edad_minima = EDAD_MINIMA_CLASIFICACION.get(clasificacion, 0)
    if edad_minima == 0:
        return True, None
    if not hash_guardado:
        return False, f'El video es clasificación {clasificacion}. Registra tu fecha de nacimiento en tu perfil antes de acceder a este contenido.'
    if not fecha_nacimiento_str:
        return False, f'El video es clasificación {clasificacion} (mínimo {edad_minima} años). Incluye "fecha_nacimiento" (YYYY-MM-DD) en la solicitud para verificar tu edad.'
    if not bcrypt.checkpw(fecha_nacimiento_str.encode('utf-8'), hash_guardado.encode('utf-8')):
        return False, 'Fecha de nacimiento incorrecta.'
    try:
        fecha_nac = datetime.datetime.strptime(fecha_nacimiento_str, '%Y-%m-%d').date()
    except ValueError:
        return False, 'Formato de fecha inválido. Use YYYY-MM-DD.'
    hoy = datetime.date.today()
    edad = hoy.year - fecha_nac.year - ((hoy.month, hoy.day) < (fecha_nac.month, fecha_nac.day))
    if edad < edad_minima:
        return False, f'Debes tener al menos {edad_minima} años para acceder a contenido clasificación {clasificacion}.'
    return True, None

transaccion_video_model = api.model('TransaccionVideo', {
    'id_video': fields.Integer(required=True),
    'fecha_nacimiento': fields.String(description='Requerida para clasificación PG-13, R o NC-17 (YYYY-MM-DD)')
})

@api.route('/api/inicio')
class Inicio(Resource):
    @token_required(roles=['CLIENTE'])
    def get(self):
        conn = get_db()
        cursor = conn.cursor()
        
        # 14 películas aleatorias
        cursor.execute("""
            SELECT TOP 14 v.id_video, v.titulo_original, v.anio, v.duracion,
                   vp.promedio, v.id_clasificacion
            FROM video v
            LEFT JOIN vw_promedio_calificaciones vp ON v.id_video = vp.id_video
            WHERE v.tipo = 'PELICULA'
            ORDER BY NEWID()
        """)
        peliculas = [{'id': r[0], 'titulo': r[1], 'anio': r[2],
                      'duracion': r[3], 'promedio': float(r[4]) if r[4] else 0,
                      'clasificacion': r[5],
                      'url_imagen': get_poster_url(r[0])}
                     for r in cursor.fetchall()]

        # Series aleatorias (una por título)
        cursor.execute("""
            SELECT TOP 5 titulo, MIN(id_serie) AS id_representativo, COUNT(id_serie) AS total_temporadas
            FROM serie
            GROUP BY titulo
            ORDER BY NEWID()
        """)
        series = [{'titulo': r[0], 'id': r[1],
             'url_imagen': get_serie_poster_url(r[1]),
             'total_temporadas': r[2]}
            for r in cursor.fetchall()]

        conn.close()
        return {'peliculas': peliculas, 'series': series}

@api.route('/api/alquilar')
class Alquilar(Resource):
    @token_required(roles=['CLIENTE'])
    @api.expect(transaccion_video_model)
    def post(self):
        data = request.json
        id_cliente = request.user.get('id')
        id_video = data.get('id_video')
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT id_video, titulo_original, id_clasificacion FROM video WHERE id_video = ?", (id_video,))
            video = cursor.fetchone()
            if not video:
                return {'error': 'Video no encontrado'}, 404

            cursor.execute("""
                SELECT id_transaccion FROM transaccion
                WHERE id_cliente = ? AND id_video = ? AND tipo = 'ALQUILER'
                AND fecha_expiracion >= GETDATE()
            """, (id_cliente, id_video))
            if cursor.fetchone():
                return {'error': 'Ya tienes este video alquilado y sigue vigente'}, 400

            cursor.execute("SELECT saldo, fecha_nacimiento_hash FROM cliente WHERE id_cliente = ?", (id_cliente,))
            row = cursor.fetchone()
            if not row or float(row[0]) < PRECIO_ALQUILER:
                return {'error': f'Saldo insuficiente. El alquiler cuesta ${PRECIO_ALQUILER}'}, 400

            ok, error_edad = verificar_restriccion_edad(video[2], data.get('fecha_nacimiento'), row[1])
            if not ok:
                return {'error': error_edad}, 403

            fecha_expiracion = (datetime.date.today() + datetime.timedelta(days=30)).isoformat()

            cursor.execute("""
                UPDATE cliente 
                SET saldo = saldo - ?, puntos = puntos + 1
                WHERE id_cliente = ?
            """, (PRECIO_ALQUILER, id_cliente))

            cursor.execute("""
                INSERT INTO transaccion (id_cliente, id_video, tipo, monto, fecha_expiracion)
                VALUES (?, ?, 'ALQUILER', ?, ?)
            """, (id_cliente, id_video, PRECIO_ALQUILER, fecha_expiracion))

            conn.commit()
            return {
                'mensaje': f'Alquiler exitoso de "{video[1]}"',
                'monto_cobrado': PRECIO_ALQUILER,
                'fecha_expiracion': fecha_expiracion,
                'puntos_ganados': 1
            }, 201
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 500
        finally:
            conn.close()


@api.route('/api/comprar')
class Comprar(Resource):
    @token_required(roles=['CLIENTE'])
    @api.expect(transaccion_video_model)
    def post(self):
        data = request.json
        id_cliente = request.user.get('id')
        id_video = data.get('id_video')
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT id_video, titulo_original, id_clasificacion FROM video WHERE id_video = ?", (id_video,))
            video = cursor.fetchone()
            if not video:
                return {'error': 'Video no encontrado'}, 404

            cursor.execute("""
                SELECT id_transaccion FROM transaccion
                WHERE id_cliente = ? AND id_video = ? AND tipo = 'COMPRA'
            """, (id_cliente, id_video))
            if cursor.fetchone():
                return {'error': 'Ya compraste este video anteriormente'}, 400

            cursor.execute("SELECT saldo, fecha_nacimiento_hash FROM cliente WHERE id_cliente = ?", (id_cliente,))
            row = cursor.fetchone()
            if not row or float(row[0]) < PRECIO_COMPRA:
                return {'error': f'Saldo insuficiente. La compra cuesta ${PRECIO_COMPRA}'}, 400

            ok, error_edad = verificar_restriccion_edad(video[2], data.get('fecha_nacimiento'), row[1])
            if not ok:
                return {'error': error_edad}, 403

            cursor.execute("""
                UPDATE cliente 
                SET saldo = saldo - ?, puntos = puntos + 2
                WHERE id_cliente = ?
            """, (PRECIO_COMPRA, id_cliente))

            cursor.execute("SELECT puntos FROM cliente WHERE id_cliente = ?", (id_cliente,))
            puntos_actuales = cursor.fetchone()[0]

            cursor.execute("""
                INSERT INTO transaccion (id_cliente, id_video, tipo, monto)
                VALUES (?, ?, 'COMPRA', ?)
            """, (id_cliente, id_video, PRECIO_COMPRA))

            conn.commit()
            respuesta = {
                'mensaje': f'Compra exitosa de "{video[1]}"',
                'monto_cobrado': PRECIO_COMPRA,
                'puntos_ganados': 2,
                'puntos_totales': puntos_actuales
            }
            if puntos_actuales >= 20:
                respuesta['aviso'] = '🎉 Tienes 20 puntos. Puedes canjear un video gratis.'
            return respuesta, 201
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 500
        finally:
            conn.close()

def get_serie_poster_url(id_serie):
    base = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'static', 'posters')
    for ext in ['jpg', 'jpeg', 'png', 'webp']:
        if os.path.exists(os.path.join(base, f'serie_{id_serie}.{ext}')):
            return f'/static/posters/serie_{id_serie}.{ext}'
    return None

@api.route('/api/series')
class ListarSeries(Resource):
    @token_required(roles=['CLIENTE'])
    def get(self):
        categoria = request.args.get('categoria', '')
        conn = get_db()
        cursor = conn.cursor()
        query = """
            SELECT s.titulo,
                   COUNT(DISTINCT s.id_serie) AS total_temporadas,
                   MIN(s.id_serie) AS id_representativo
            FROM serie s
            LEFT JOIN serie_video sv ON s.id_serie = sv.id_serie
            LEFT JOIN video v ON sv.id_video = v.id_video
            LEFT JOIN video_categoria vc ON v.id_video = vc.id_video
            LEFT JOIN categoria c ON vc.id_categoria = c.id_categoria
            WHERE 1=1
        """
        params = []
        if categoria:
            query += " AND c.nombre = ?"
            params.append(categoria.upper())
        query += " GROUP BY s.titulo"
        cursor.execute(query, params)
        rows = cursor.fetchall()
        conn.close()
        series = [{
            'titulo': r[0],
            'total_temporadas': r[1],
            'id': r[2],
            'url_imagen': get_serie_poster_url(r[2])
        } for r in rows]
        return {'series': series}
    
@api.route('/api/series/<string:titulo>/temporadas')
class TemporadasSerie(Resource):
    @token_required(roles=['CLIENTE'])
    def get(self, titulo):
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT id_serie, temporada
            FROM serie
            WHERE titulo = ?
            ORDER BY temporada
        """, (titulo,))
        rows = cursor.fetchall()
        conn.close()
        return {'temporadas': [{'id': r[0], 'temporada': r[1]} for r in rows]}

@api.route('/api/series/<int:id_serie>/episodios')
class EpisodiosSerie(Resource):
    @token_required(roles=['CLIENTE'])
    def get(self, id_serie):
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT v.id_video, v.titulo_original, v.duracion
            FROM serie_video sv
            JOIN video v ON sv.id_video = v.id_video
            WHERE sv.id_serie = ?
            ORDER BY v.id_video
        """, (id_serie,))
        rows = cursor.fetchall()
        conn.close()
        return {'episodios': [{'id': r[0], 'titulo': r[1], 'duracion': r[2]} for r in rows]}

@api.route('/api/video/<int:id_video>/idiomas')
class VideoIdiomas(Resource):
    @token_required(roles=['CLIENTE'])
    def get(self, id_video):
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT i.lenguaje, vi.tipo
            FROM video_idioma vi
            JOIN idioma i ON vi.id_idioma = i.id_idioma
            WHERE vi.id_video = ?
            ORDER BY vi.tipo, i.lenguaje
        """, (id_video,))
        rows = cursor.fetchall()
        conn.close()
        return {'idiomas': [{'lenguaje': r[0], 'tipo': r[1]} for r in rows]}

@api.route('/api/video/<int:id_video>/personas-idiomas')
class VideoPersonasIdiomas(Resource):
    @token_required(roles=['CLIENTE'])
    def get(self, id_video):
        conn = get_db()
        cursor = conn.cursor()

        cursor.execute("SELECT TOP 1 id_serie FROM serie_video WHERE id_video = ?", (id_video,))
        row_serie = cursor.fetchone()
        es_serie = row_serie is not None

        if es_serie:
            cursor.execute("""
                SELECT p.nombre, sp.tipo
                FROM serie_persona sp
                JOIN persona p ON sp.id_persona = p.id_persona
                WHERE sp.id_serie = ?
                ORDER BY sp.tipo, p.nombre
            """, (row_serie[0],))
        else:
            cursor.execute("""
                SELECT p.nombre, vp.rol
                FROM video_persona vp
                JOIN persona p ON vp.id_persona = p.id_persona
                WHERE vp.id_video = ?
                ORDER BY vp.rol, p.nombre
            """, (id_video,))
        personas = cursor.fetchall()

        cursor.execute("""
            SELECT i.lenguaje, vi.tipo
            FROM video_idioma vi
            JOIN idioma i ON vi.id_idioma = i.id_idioma
            WHERE vi.id_video = ?
            ORDER BY vi.tipo, i.lenguaje
        """, (id_video,))
        idiomas = cursor.fetchall()
        conn.close()
        return {
            'es_serie': es_serie,
            'directores': [r[0] for r in personas if r[1] == 'DIRECTOR'],
            'actores': [r[0] for r in personas if r[1] == 'ACTOR'],
            'productores': [r[0] for r in personas if r[1] == 'PRODUCTOR'],
            'idiomas': [{'lenguaje': r[0], 'tipo': r[1]} for r in idiomas]
        }

@api.route('/api/mis-videos')
class MisVideos(Resource):
    @token_required(roles=['CLIENTE'])
    def get(self):
        id_cliente = request.user.get('id')
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT t.id_transaccion, v.titulo_original, t.tipo, t.fecha,
                   t.monto, t.fecha_expiracion, v.id_video, v.duracion, v.descripcion,
                   STRING_AGG(c.nombre, ', ') AS categorias, v.url_descarga
            FROM transaccion t
            JOIN video v ON t.id_video = v.id_video
            LEFT JOIN video_categoria vc ON v.id_video = vc.id_video
            LEFT JOIN categoria c ON vc.id_categoria = c.id_categoria
            WHERE t.id_cliente = ?
            GROUP BY t.id_transaccion, v.titulo_original, t.tipo, t.fecha,
                     t.monto, t.fecha_expiracion, v.id_video, v.duracion, v.descripcion, v.url_descarga
            ORDER BY t.fecha DESC
        """, (id_cliente,))
        rows = cursor.fetchall()
        conn.close()
        videos = []
        for r in rows:
            item = {
                'id_transaccion': r[0],
                'titulo': r[1],
                'tipo': r[2],
                'fecha': r[3].strftime('%Y-%m-%d') if r[3] else None,
                'id_video': r[6],
                'duracion': r[7],
                'descripcion': r[8],
                'categorias': r[9] if r[9] else 'Sin categoría',
                'url_imagen': get_poster_url(r[6]),
                'url_descarga': r[10]
            }
            if r[2] == 'ALQUILER':
                item['fecha_expiracion'] = str(r[5]) if r[5] else None
                item['activo'] = r[5] >= datetime.date.today() if r[5] else False
            videos.append(item)
        return {'videos': videos}


@api.route('/api/mis-puntos')
class MisPuntos(Resource):
    @token_required(roles=['CLIENTE'])
    def get(self):
        id_cliente = request.user.get('id')
        conn = get_db()
        cursor = conn.cursor()

        cursor.execute("SELECT puntos FROM cliente WHERE id_cliente = ?", (id_cliente,))
        row = cursor.fetchone()
        puntos = int(row[0]) if row else 0

        cursor.execute("""
            SELECT v.titulo_original, t.tipo, t.fecha,
                   CASE t.tipo WHEN 'COMPRA' THEN 2 ELSE 1 END AS pts
            FROM transaccion t
            JOIN video v ON t.id_video = v.id_video
            WHERE t.id_cliente = ?
            ORDER BY t.fecha DESC
        """, (id_cliente,))
        historial = [{
            'titulo': r[0],
            'tipo': r[1],
            'fecha': r[2].strftime('%Y-%m-%d') if r[2] else None,
            'puntos': int(r[3])
        } for r in cursor.fetchall()]

        cursor.execute("SELECT COUNT(*) FROM cliente WHERE id_referido = ?", (id_cliente,))
        total_referidos = int(cursor.fetchone()[0])

        conn.close()
        return {
            'puntos': puntos,
            'historial': historial,
            'total_referidos': total_referidos
        }

@api.route('/api/categorias-cliente')
class CategoriasCliente(Resource):
    @token_required(roles=['CLIENTE'])
    def get(self):
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT id_categoria, nombre FROM categoria ORDER BY nombre")
        rows = cursor.fetchall()
        conn.close()
        return {'categorias': [{'id': r[0], 'nombre': r[1]} for r in rows]}

@api.route('/api/canjear-puntos')
class CanjearPuntos(Resource):
    @token_required(roles=['CLIENTE'])
    def post(self):
        id_cliente = request.user.get('id')
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT puntos FROM cliente WHERE id_cliente = ?", (id_cliente,))
        row = cursor.fetchone()
        if not row:
            return {'error': 'Cliente no encontrado'}, 404
        puntos = row[0]
        if puntos >= 20:
            cursor.execute("UPDATE cliente SET puntos = puntos - 20 WHERE id_cliente = ?", (id_cliente,))
            conn.commit()
            conn.close()
            return {'mensaje': '¡Felicidades! Has canjeado un video gratis.'}
        conn.close()
        return {'error': 'Necesitas 20 puntos para canjear un video gratis'}, 400

@api.route('/api/referidos')
class Referidos(Resource):
    @token_required(roles=['CLIENTE'])
    def get(self):
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute(
            "SELECT nombre_amigo, email_amigo, estado FROM referido WHERE id_cliente = ?",
            (request.user.get('id'),)
        )
        rows = cursor.fetchall()
        conn.close()
        return {'referidos': [{'nombre': r[0], 'email': r[1], 'estado': r[2]} for r in rows]}

# ============================================
# ENDPOINTS ADMINISTRADOR
# ============================================

@api.route('/admin/categorias')
class AdminCategorias(Resource):
    @token_required(roles=['ADMINISTRADOR'])
    def get(self):
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT id_categoria, nombre FROM categoria")
        rows = cursor.fetchall()
        conn.close()
        return {'categorias': [{'id': r[0], 'nombre': r[1]} for r in rows]}
    
    @token_required(roles=['ADMINISTRADOR'])
    @api.expect(categoria_model)
    def post(self):
        data = request.json
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("INSERT INTO categoria (nombre) VALUES (?)", (data['nombre'].upper(),))
            conn.commit()
            return {'mensaje': 'Categoría creada exitosamente'}, 201
        except Exception as e:
            return {'error': str(e)}, 400
        finally:
            conn.close()

@api.route('/admin/personas')
class AdminPersonas(Resource):
    @token_required(roles=['ADMINISTRADOR'])
    def get(self):
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT id_persona, nombre, fecha_nacimiento FROM persona ORDER BY nombre")
        rows = cursor.fetchall()
        conn.close()
        return {'personas': [{'id': r[0], 'nombre': r[1],
                               'fecha_nacimiento': str(r[2]) if r[2] else None} for r in rows]}

    @token_required(roles=['ADMINISTRADOR'])
    @api.expect(persona_model)
    def post(self):
        data = request.json
        fecha_nac = data.get('fecha_nacimiento') or None
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute(
                "INSERT INTO persona (nombre, fecha_nacimiento) VALUES (?, ?)",
                (data['nombre'], fecha_nac)
            )
            cursor.execute("SELECT @@IDENTITY")
            id_persona = cursor.fetchone()[0]
            conn.commit()
            return {'id_persona': id_persona, 'mensaje': 'Persona creada exitosamente'}, 201
        except Exception as e:
            return {'error': str(e)}, 400
        finally:
            conn.close()

@api.route('/admin/videos')
class AdminVideos(Resource):
    @token_required(roles=['ADMINISTRADOR'])
    def get(self):
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT v.id_video, v.isan, v.titulo_original, v.anio, v.duracion,
                   v.descripcion,
                   STRING_AGG(CASE WHEN vper.rol = 'DIRECTOR' THEN per.nombre END, ', ') AS directores,
                   STRING_AGG(CASE WHEN vper.rol = 'ACTOR' THEN per.nombre END, ', ') AS actores
            FROM video v
            LEFT JOIN video_persona vper ON v.id_video = vper.id_video
            LEFT JOIN persona per ON vper.id_persona = per.id_persona
            GROUP BY v.id_video, v.isan, v.titulo_original, v.anio, v.duracion, v.descripcion
        """)
        rows = cursor.fetchall()
        conn.close()
        return {'videos': [{
            'id': r[0],
            'isan': r[1],
            'titulo': r[2],
            'anio': r[3],
            'duracion': r[4],
            'descripcion': r[5],
            'directores': r[6].split(', ') if r[6] else [],
            'actores': r[7].split(', ') if r[7] else []
        } for r in rows]}

    @token_required(roles=['ADMINISTRADOR'])
    @api.expect(video_model)
    def post(self):
        data = request.json
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT id_video FROM video WHERE isan = ?", (data['isan'],))
            if cursor.fetchone():
                return {'error': 'El ISAN ya existe'}, 400

            # Insertar directo en video, sin tabla pelicula
            cursor.execute("""
                INSERT INTO video (isan, titulo_original, anio, duracion, tipo, descripcion, id_clasificacion)
                VALUES (?, ?, ?, ?, 'PELICULA', ?, ?)
            """, (data['isan'], data['titulo_original'], data['anio'], data['duracion'],
                  data.get('descripcion'), data.get('id_clasificacion')))

            cursor.execute("SELECT @@IDENTITY")
            id_video = int(cursor.fetchone()[0])

            # Vincular categorías
            for categoria in data.get('categorias', [])[:3]:
                cursor.execute("SELECT id_categoria FROM categoria WHERE nombre = ?", (categoria.upper(),))
                cat = cursor.fetchone()
                if cat:
                    cursor.execute("INSERT INTO video_categoria (id_video, id_categoria) VALUES (?, ?)",
                                   (id_video, cat[0]))

            # Vincular personas
            for persona in data.get('personas', []):
                cursor.execute("SELECT id_persona FROM persona WHERE nombre = ?", (persona['nombre'],))
                row = cursor.fetchone()
                if row:
                    id_persona = row[0]
                else:
                    cursor.execute("INSERT INTO persona (nombre) VALUES (?)", (persona['nombre'],))
                    cursor.execute("SELECT @@IDENTITY")
                    id_persona = int(cursor.fetchone()[0])
                cursor.execute("INSERT INTO video_persona (id_video, id_persona, rol) VALUES (?, ?, ?)",
                               (id_video, id_persona, persona['rol']))

            conn.commit()
            return {'mensaje': 'Video creado exitosamente', 'id_video': id_video}, 201
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 500
        finally:
            conn.close()

@api.route('/admin/video/<int:id_video>')
class AdminVideoDetail(Resource):
    @token_required(roles=['ADMINISTRADOR'])
    @api.expect(video_model)
    def put(self, id_video):
        data = request.json
        conn = get_db()
        cursor = conn.cursor()
        try:
            # Todo va directo en video, sin tabla pelicula
            cursor.execute("""
                UPDATE video 
                SET isan=?, titulo_original=?, anio=?, duracion=?, 
                    descripcion=?, id_clasificacion=?
                WHERE id_video=?
            """, (data['isan'], data['titulo_original'], data['anio'], data['duracion'],
                  data.get('descripcion'), data.get('id_clasificacion'), id_video))

            # Actualizar personas si vienen en el request
            if data.get('personas'):
                cursor.execute("DELETE FROM video_persona WHERE id_video = ?", (id_video,))
                for persona in data.get('personas', []):
                    cursor.execute("SELECT id_persona FROM persona WHERE nombre = ?", (persona['nombre'],))
                    row = cursor.fetchone()
                    if row:
                        id_persona = row[0]
                    else:
                        cursor.execute("INSERT INTO persona (nombre) VALUES (?)", (persona['nombre'],))
                        cursor.execute("SELECT @@IDENTITY")
                        id_persona = int(cursor.fetchone()[0])
                    cursor.execute("INSERT INTO video_persona (id_video, id_persona, rol) VALUES (?, ?, ?)",
                                   (id_video, id_persona, persona['rol']))

            # Actualizar categorías si vienen en el request
            if data.get('categorias'):
                cursor.execute("DELETE FROM video_categoria WHERE id_video = ?", (id_video,))
                for categoria in data.get('categorias', [])[:3]:
                    cursor.execute("SELECT id_categoria FROM categoria WHERE nombre = ?", (categoria.upper(),))
                    cat = cursor.fetchone()
                    if cat:
                        cursor.execute("INSERT INTO video_categoria (id_video, id_categoria) VALUES (?, ?)",
                                       (id_video, cat[0]))

            conn.commit()
            return {'mensaje': 'Video actualizado exitosamente'}
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 500
        finally:
            conn.close()

    @token_required(roles=['ADMINISTRADOR'])
    def delete(self, id_video):
        conn = get_db()
        cursor = conn.cursor()
        try:
            # Las FKs con CASCADE se eliminan automáticamente,
            # pero por claridad se hace explícito
            cursor.execute("DELETE FROM calificacion WHERE id_video = ?", (id_video,))
            cursor.execute("DELETE FROM video_persona WHERE id_video = ?", (id_video,))
            cursor.execute("DELETE FROM video_categoria WHERE id_video = ?", (id_video,))
            cursor.execute("DELETE FROM video_idioma WHERE id_video = ?", (id_video,))
            cursor.execute("DELETE FROM video WHERE id_video = ?", (id_video,))
            conn.commit()
            return {'mensaje': 'Video eliminado exitosamente'}
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 500
        finally:
            conn.close()

@api.route('/admin/video/<int:id_video>/categorias')
class AdminVideoCategorias(Resource):
    @token_required(roles=['ADMINISTRADOR'])
    @api.expect(api.model('AsignarCategorias', {'categorias': fields.List(fields.String)}))
    def post(self, id_video):
        data = request.json
        conn = get_db()
        cursor = conn.cursor()
        try:
            categorias = data.get('categorias', [])
            if len(categorias) > 3:
                return {'error': 'Un video puede tener máximo 3 categorías'}, 400

            cursor.execute("DELETE FROM video_categoria WHERE id_video = ?", (id_video,))
            for categoria in categorias:
                cursor.execute("SELECT id_categoria FROM categoria WHERE nombre = ?", (categoria.upper(),))
                cat = cursor.fetchone()
                if cat:
                    # Sin columna orden, fue eliminada
                    cursor.execute("INSERT INTO video_categoria (id_video, id_categoria) VALUES (?, ?)",
                                   (id_video, cat[0]))
                else:
                    return {'error': f'Categoría "{categoria}" no existe'}, 400

            conn.commit()
            return {'mensaje': 'Categorías asignadas exitosamente'}
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 500
        finally:
            conn.close()


@api.route('/admin/video/<int:id_video>/personas')
class AdminVideoPersonas(Resource):
    @token_required(roles=['ADMINISTRADOR'])
    @api.expect(api.model('AsignarPersonas', {'personas': fields.List(fields.Nested(api.model('PersonaAsignar', {'nombre': fields.String, 'rol': fields.String})))}))
    def post(self, id_video):
        data = request.json
        conn = get_db()
        cursor = conn.cursor()
        try:
            # Limpiar personas anteriores antes de asignar nuevas
            cursor.execute("DELETE FROM video_persona WHERE id_video = ?", (id_video,))

            for persona in data.get('personas', []):
                if persona['rol'] not in ('ACTOR', 'DIRECTOR', 'PRODUCTOR'):
                    return {'error': f'Rol "{persona["rol"]}" no válido'}, 400

                cursor.execute("SELECT id_persona FROM persona WHERE nombre = ?", (persona['nombre'],))
                row = cursor.fetchone()
                if row:
                    id_persona = row[0]
                else:
                    cursor.execute("INSERT INTO persona (nombre) VALUES (?)", (persona['nombre'],))
                    cursor.execute("SELECT @@IDENTITY")
                    id_persona = int(cursor.fetchone()[0])

                cursor.execute("INSERT INTO video_persona (id_video, id_persona, rol) VALUES (?, ?, ?)",
                               (id_video, id_persona, persona['rol']))

            conn.commit()
            return {'mensaje': 'Personas asignadas exitosamente'}
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 500
        finally:
            conn.close()

@api.route('/admin/series')
class AdminSeries(Resource):
    @token_required(roles=['ADMINISTRADOR'])
    def get(self):
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT s.id_serie, s.titulo, s.temporada,
                   COUNT(sv.id_video) AS total_episodios
            FROM serie s
            LEFT JOIN serie_video sv ON s.id_serie = sv.id_serie
            GROUP BY s.id_serie, s.titulo, s.temporada
            ORDER BY s.titulo, s.temporada
        """)
        rows = cursor.fetchall()
        conn.close()
        return {'series': [{'id': r[0], 'titulo': r[1], 'temporada': r[2],
                             'total_episodios': r[3]} for r in rows]}

@api.route('/admin/series/<int:id_serie>/personas')
class AdminSeriePersonas(Resource):
    @token_required(roles=['ADMINISTRADOR'])
    def get(self, id_serie):
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT p.id_persona, p.nombre, sp.tipo
            FROM serie_persona sp
            JOIN persona p ON sp.id_persona = p.id_persona
            WHERE sp.id_serie = ?
            ORDER BY sp.tipo, p.nombre
        """, (id_serie,))
        rows = cursor.fetchall()
        conn.close()
        return {'personas': [{'id': r[0], 'nombre': r[1], 'tipo': r[2]} for r in rows]}

    @token_required(roles=['ADMINISTRADOR'])
    @api.expect(api.model('AsignarPersonasSerie', {'personas': fields.List(fields.Nested(api.model('PersonaSerieAsignar', {'nombre': fields.String, 'tipo': fields.String(enum=['ACTOR', 'DIRECTOR', 'PRODUCTOR', 'INVITADO'])})))}))
    def post(self, id_serie):
        data = request.json
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("DELETE FROM serie_persona WHERE id_serie = ?", (id_serie,))
            for persona in data.get('personas', []):
                if persona['tipo'] not in ('ACTOR', 'DIRECTOR', 'PRODUCTOR', 'INVITADO'):
                    return {'error': f'Tipo "{persona["tipo"]}" no válido. Use: ACTOR, DIRECTOR, PRODUCTOR o INVITADO'}, 400
                cursor.execute("SELECT id_persona FROM persona WHERE nombre = ?", (persona['nombre'],))
                row = cursor.fetchone()
                if row:
                    id_persona = row[0]
                else:
                    cursor.execute("INSERT INTO persona (nombre) VALUES (?)", (persona['nombre'],))
                    cursor.execute("SELECT @@IDENTITY")
                    id_persona = int(cursor.fetchone()[0])
                cursor.execute("INSERT INTO serie_persona (id_serie, id_persona, tipo) VALUES (?, ?, ?)",
                               (id_serie, id_persona, persona['tipo']))
            conn.commit()
            return {'mensaje': 'Personas asignadas a la serie exitosamente'}
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 500
        finally:
            conn.close()

# ============================================
# ADMIN — IDIOMAS
# ============================================

@api.route('/admin/idiomas')
class AdminIdiomas(Resource):
    @token_required(roles=['ADMINISTRADOR'])
    def get(self):
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT id_idioma, lenguaje FROM idioma ORDER BY lenguaje")
        rows = cursor.fetchall()
        conn.close()
        return {'idiomas': [{'id': r[0], 'lenguaje': r[1]} for r in rows]}

    @token_required(roles=['ADMINISTRADOR'])
    @api.expect(api.model('Idioma', {'lenguaje': fields.String(required=True)}))
    def post(self):
        data = request.json
        lenguaje = data.get('lenguaje', '').strip()
        if not lenguaje:
            return {'error': 'El lenguaje es requerido'}, 400
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT id_idioma FROM idioma WHERE lenguaje = ?", (lenguaje,))
            if cursor.fetchone():
                return {'error': 'El idioma ya existe'}, 400
            cursor.execute("INSERT INTO idioma (lenguaje) VALUES (?)", (lenguaje,))
            cursor.execute("SELECT @@IDENTITY")
            id_idioma = int(cursor.fetchone()[0])
            conn.commit()
            return {'mensaje': 'Idioma creado exitosamente', 'id_idioma': id_idioma}, 201
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 400
        finally:
            conn.close()

@api.route('/admin/video/<int:id_video>/idiomas')
class AdminVideoIdiomas(Resource):
    @token_required(roles=['ADMINISTRADOR'])
    def get(self, id_video):
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT i.id_idioma, i.lenguaje, vi.tipo
            FROM video_idioma vi
            JOIN idioma i ON vi.id_idioma = i.id_idioma
            WHERE vi.id_video = ?
            ORDER BY vi.tipo, i.lenguaje
        """, (id_video,))
        rows = cursor.fetchall()
        conn.close()
        return {'idiomas': [{'id': r[0], 'lenguaje': r[1], 'tipo': r[2]} for r in rows]}

    @token_required(roles=['ADMINISTRADOR'])
    @api.expect(api.model('AsignarIdiomas', {
        'idiomas': fields.List(fields.Nested(api.model('IdiomaAsignar', {
            'id_idioma': fields.Integer(required=True),
            'tipo': fields.String(required=True, enum=['ORIGINAL', 'SUBTITULO', 'DOBLAJE'])
        })))
    }))
    def post(self, id_video):
        data = request.json
        idiomas = data.get('idiomas', [])
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT id_video FROM video WHERE id_video = ?", (id_video,))
            if not cursor.fetchone():
                return {'error': 'Video no encontrado'}, 404
            cursor.execute("DELETE FROM video_idioma WHERE id_video = ?", (id_video,))
            for item in idiomas:
                if item['tipo'] not in ('ORIGINAL', 'SUBTITULO', 'DOBLAJE'):
                    return {'error': f'Tipo "{item["tipo"]}" no válido. Use ORIGINAL, SUBTITULO o DOBLAJE'}, 400
                cursor.execute("SELECT id_idioma FROM idioma WHERE id_idioma = ?", (item['id_idioma'],))
                if not cursor.fetchone():
                    return {'error': f'Idioma {item["id_idioma"]} no encontrado'}, 404
                cursor.execute(
                    "INSERT INTO video_idioma (id_video, id_idioma, tipo) VALUES (?, ?, ?)",
                    (id_video, item['id_idioma'], item['tipo'])
                )
            conn.commit()
            return {'mensaje': 'Idiomas asignados exitosamente'}
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 500
        finally:
            conn.close()

# ============================================
# ADMIN — CLASIFICACIONES
# ============================================

@api.route('/admin/clasificaciones')
class AdminClasificaciones(Resource):
    @token_required(roles=['ADMINISTRADOR'])
    def get(self):
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT tipo, descripcion FROM clasificacion ORDER BY tipo")
            rows = cursor.fetchall()
            conn.close()
            return {'clasificaciones': [{'tipo': r[0], 'descripcion': r[1]} for r in rows]}
        except Exception as e:
            conn.close()
            return {'error': str(e)}, 500

    @token_required(roles=['ADMINISTRADOR'])
    @api.expect(api.model('Clasificacion', {
        'tipo': fields.String(required=True, enum=['G', 'PG', 'PG-13', 'R', 'NC-17']),
        'descripcion': fields.String(required=True)
    }))
    def post(self):
        data = request.json
        tipo = data.get('tipo', '').upper()
        descripcion = data.get('descripcion', '').strip()
        if tipo not in ('G', 'PG', 'PG-13', 'R', 'NC-17'):
            return {'error': 'Tipo debe ser G, PG, PG-13, R o NC-17'}, 400
        if not descripcion:
            return {'error': 'La descripción es requerida'}, 400
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT tipo FROM clasificacion WHERE tipo = ?", (tipo,))
            if cursor.fetchone():
                return {'error': f'La clasificación {tipo} ya existe'}, 400
            cursor.execute("INSERT INTO clasificacion (tipo, descripcion) VALUES (?, ?)", (tipo, descripcion))
            conn.commit()
            return {'mensaje': f'Clasificación {tipo} creada exitosamente'}, 201
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 400
        finally:
            conn.close()

@api.route('/admin/clasificacion/<string:tipo>')
class AdminClasificacionDetalle(Resource):
    @token_required(roles=['ADMINISTRADOR'])
    @api.expect(api.model('ActualizarClasificacion', {'descripcion': fields.String(required=True)}))
    def put(self, tipo):
        data = request.json
        descripcion = data.get('descripcion', '').strip()
        if not descripcion:
            return {'error': 'La descripción es requerida'}, 400
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("UPDATE clasificacion SET descripcion = ? WHERE tipo = ?", (descripcion, tipo.upper()))
            if cursor.rowcount == 0:
                return {'error': f'Clasificación {tipo} no encontrada'}, 404
            conn.commit()
            return {'mensaje': 'Clasificación actualizada exitosamente'}
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 400
        finally:
            conn.close()

@api.route('/admin/video/<int:id_video>/clasificacion')
class AdminVideoClasificacion(Resource):
    @token_required(roles=['ADMINISTRADOR'])
    @api.expect(api.model('AsignarClasificacion', {
        'tipo': fields.String(required=True, enum=['G', 'PG', 'PG-13', 'R', 'NC-17'])
    }))
    def put(self, id_video):
        data = request.json
        tipo = data.get('tipo', '').upper()
        if tipo not in ('G', 'PG', 'PG-13', 'R', 'NC-17'):
            return {'error': 'Tipo debe ser G, PG, PG-13, R o NC-17'}, 400
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("UPDATE video SET id_clasificacion = ? WHERE id_video = ?", (tipo, id_video))
            if cursor.rowcount == 0:
                return {'error': 'Video no encontrado'}, 404
            conn.commit()
            return {'mensaje': f'Clasificación {tipo} asignada al video exitosamente'}
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 400
        finally:
            conn.close()

# ============================================
# ADMIN — GESTIÓN DE USUARIOS
# ============================================

usuario_admin_model = api.model('UsuarioAdmin', {
    'nombre_completo': fields.String(required=True),
    'documento': fields.String(required=True),
    'nombre_usuario': fields.String(required=True, description='Usado para iniciar sesión'),
    'password': fields.String(required=True),
    'rol': fields.String(required=True, enum=['ADMINISTRADOR', 'GERENTE'])
})

@api.route('/admin/usuarios')
class AdminUsuarios(Resource):
    @token_required(roles=['ADMINISTRADOR'])
    def get(self):
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT id_usuario, nombre_completo, documento, nombre_usuario, rol, fecha_creacion
            FROM usuario_sistema
            ORDER BY rol, nombre_completo
        """)
        rows = cursor.fetchall()
        conn.close()
        return {'usuarios': [
            {'id': r[0], 'nombre_completo': r[1], 'documento': r[2],
             'nombre_usuario': r[3], 'rol': r[4],
             'fecha_creacion': str(r[5]) if r[5] else None}
            for r in rows
        ]}

    @token_required(roles=['ADMINISTRADOR'])
    @api.expect(usuario_admin_model)
    def post(self):
        data = request.json
        rol = data.get('rol', '').upper()
        if rol not in ('ADMINISTRADOR', 'GERENTE'):
            return {'error': 'El rol debe ser ADMINISTRADOR o GERENTE'}, 400
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT id_usuario FROM usuario_sistema WHERE nombre_usuario = ?",
                           (data['nombre_usuario'],))
            if cursor.fetchone():
                return {'error': 'El nombre de usuario ya existe'}, 400
            cursor.execute("SELECT id_usuario FROM usuario_sistema WHERE documento = ?",
                           (data['documento'],))
            if cursor.fetchone():
                return {'error': 'El documento ya está registrado'}, 400
            hashed = bcrypt.hashpw(data['password'].encode('utf-8'), bcrypt.gensalt())
            cursor.execute("""
                INSERT INTO usuario_sistema (nombre_completo, documento, nombre_usuario, password, rol)
                VALUES (?, ?, ?, ?, ?)
            """, (data['nombre_completo'], data['documento'],
                  data['nombre_usuario'], hashed.decode('utf-8'), rol))
            cursor.execute("SELECT @@IDENTITY")
            id_nuevo = int(cursor.fetchone()[0])
            conn.commit()
            return {'mensaje': f'Usuario {rol} creado exitosamente', 'id': id_nuevo}, 201
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 400
        finally:
            conn.close()

@api.route('/admin/usuario/<int:id_usuario>')
class AdminUsuarioDetalle(Resource):
    @token_required(roles=['ADMINISTRADOR'])
    @api.expect(api.model('ActualizarUsuario', {
        'nombre_completo': fields.String,
        'documento': fields.String,
        'nombre_usuario': fields.String
    }))
    def put(self, id_usuario):
        data = request.json
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("""
                UPDATE usuario_sistema
                SET nombre_completo=?, documento=?, nombre_usuario=?
                WHERE id_usuario=?
            """, (data.get('nombre_completo'), data.get('documento'),
                  data.get('nombre_usuario'), id_usuario))
            if cursor.rowcount == 0:
                return {'error': 'Usuario no encontrado'}, 404
            conn.commit()
            return {'mensaje': 'Usuario actualizado exitosamente'}
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 400
        finally:
            conn.close()

    @token_required(roles=['ADMINISTRADOR'])
    def delete(self, id_usuario):
        if id_usuario == request.user.get('id'):
            return {'error': 'No puedes eliminar tu propio usuario'}, 400
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("DELETE FROM usuario_sistema WHERE id_usuario=?", (id_usuario,))
            if cursor.rowcount == 0:
                return {'error': 'Usuario no encontrado'}, 404
            conn.commit()
            return {'mensaje': 'Usuario eliminado exitosamente'}
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 400
        finally:
            conn.close()

@api.route('/admin/usuario/<int:id_usuario>/rol')
class AdminUsuarioRol(Resource):
    @token_required(roles=['ADMINISTRADOR'])
    @api.expect(api.model('CambiarRol', {
        'rol': fields.String(required=True, enum=['ADMINISTRADOR', 'GERENTE'])
    }))
    def put(self, id_usuario):
        rol = request.json.get('rol', '').upper()
        if rol not in ('ADMINISTRADOR', 'GERENTE'):
            return {'error': 'El rol debe ser ADMINISTRADOR o GERENTE'}, 400
        if id_usuario == request.user.get('id'):
            return {'error': 'No puedes cambiar tu propio rol'}, 400
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("UPDATE usuario_sistema SET rol=? WHERE id_usuario=?",
                           (rol, id_usuario))
            if cursor.rowcount == 0:
                return {'error': 'Usuario no encontrado'}, 404
            conn.commit()
            return {'mensaje': f'Rol actualizado a {rol} exitosamente'}
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 400
        finally:
            conn.close()

@api.route('/admin/usuario/<int:id_usuario>/reset-password')
class AdminResetPassword(Resource):
    @token_required(roles=['ADMINISTRADOR'])
    @api.expect(api.model('ResetPassword', {'nueva_password': fields.String(required=True)}))
    def put(self, id_usuario):
        nueva = request.json.get('nueva_password', '').strip()
        if not nueva or len(nueva) < 6:
            return {'error': 'La contraseña debe tener al menos 6 caracteres'}, 400
        conn = get_db()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT id_usuario FROM usuario_sistema WHERE id_usuario=?",
                           (id_usuario,))
            if not cursor.fetchone():
                return {'error': 'Usuario no encontrado'}, 404
            hashed = bcrypt.hashpw(nueva.encode('utf-8'), bcrypt.gensalt())
            cursor.execute("UPDATE usuario_sistema SET password=? WHERE id_usuario=?",
                           (hashed.decode('utf-8'), id_usuario))
            conn.commit()
            return {'mensaje': 'Contraseña restablecida exitosamente'}
        except Exception as e:
            conn.rollback()
            return {'error': str(e)}, 400
        finally:
            conn.close()

# ============================================
# ENDPOINTS GERENTE
# ============================================

@api.route('/gerente/reportes')
class Reportes(Resource):
    @token_required(roles=['GERENTE'])
    def get(self):
        conn = get_db()
        cursor = conn.cursor()

        # Reporte: mejor calificados
        cursor.execute("""
            SELECT TOP 5 titulo_original, promedio
            FROM vw_promedio_calificaciones
            WHERE promedio > 0
            ORDER BY promedio DESC
        """)
        mejor_calificados = [{'titulo': r[0], 'promedio': float(r[1])} for r in cursor.fetchall()]

        # Reporte: peor calificados
        cursor.execute("""
            SELECT TOP 5 titulo_original, promedio
            FROM vw_promedio_calificaciones
            WHERE promedio > 0
            ORDER BY promedio ASC
        """)
        peor_calificados = [{'titulo': r[0], 'promedio': float(r[1])} for r in cursor.fetchall()]

        # Reporte: actividad de clientes (transacciones por tipo)
        cursor.execute("""
            SELECT tipo, COUNT(*) AS cantidad, SUM(monto) AS total_ingresos
            FROM transaccion
            GROUP BY tipo
        """)
        transacciones_por_tipo = [
            {'tipo': r[0], 'cantidad': r[1], 'total_ingresos': float(r[2])}
            for r in cursor.fetchall()
        ]

        cursor.execute("SELECT COUNT(*) FROM calificacion WHERE fecha >= DATEADD(day, -30, GETDATE())")
        calificaciones_mes = cursor.fetchone()[0]

        cursor.execute("SELECT COUNT(*) FROM cliente WHERE rol = 'CLIENTE'")
        total_clientes = cursor.fetchone()[0]

        # Reporte: crecimiento de catálogo (videos por año de producción)
        cursor.execute("""
            SELECT anio, COUNT(*) AS cantidad
            FROM video
            WHERE anio IS NOT NULL
            GROUP BY anio
            ORDER BY anio DESC
        """)
        crecimiento_catalogo = [{'anio': r[0], 'videos': r[1]} for r in cursor.fetchall()]

        # Reporte: categorías más populares (por cantidad de transacciones)
        cursor.execute("""
            SELECT TOP 6 c.nombre, COUNT(t.id_transaccion) AS total_transacciones
            FROM categoria c
            JOIN video_categoria vc ON c.id_categoria = vc.id_categoria
            JOIN transaccion t ON vc.id_video = t.id_video
            GROUP BY c.nombre
            ORDER BY total_transacciones DESC
        """)
        categorias_populares = [{'categoria': r[0], 'transacciones': r[1]} for r in cursor.fetchall()]

        conn.close()
        return {
            'mejor_calificados': mejor_calificados,
            'peor_calificados': peor_calificados,
            'transacciones_por_tipo': transacciones_por_tipo,
            'calificaciones_ultimo_mes': calificaciones_mes,
            'total_clientes': total_clientes,
            'crecimiento_catalogo': crecimiento_catalogo,
            'categorias_populares': categorias_populares
        }

@api.route('/gerente/reportes/exportar')
class ExportarReporte(Resource):
    @token_required(roles=['GERENTE'])
    def get(self):
        formato = request.args.get('formato', 'json')
        conn = get_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT TOP 5 titulo_original, promedio FROM vw_promedio_calificaciones
            WHERE promedio > 0 ORDER BY promedio DESC
        """)
        mejor = cursor.fetchall()

        cursor.execute("""
            SELECT TOP 5 titulo_original, promedio FROM vw_promedio_calificaciones
            WHERE promedio > 0 ORDER BY promedio ASC
        """)
        peor = cursor.fetchall()

        cursor.execute("""
            SELECT tipo, COUNT(*) AS cantidad, SUM(monto) AS total_ingresos
            FROM transaccion GROUP BY tipo
        """)
        transacciones = cursor.fetchall()

        cursor.execute("""
            SELECT TOP 6 c.nombre, COUNT(t.id_transaccion) AS total
            FROM categoria c
            JOIN video_categoria vc ON c.id_categoria = vc.id_categoria
            JOIN transaccion t ON vc.id_video = t.id_video
            GROUP BY c.nombre ORDER BY total DESC
        """)
        categorias = cursor.fetchall()

        conn.close()

        if formato == 'pdf':
            buffer = BytesIO()
            c = canvas.Canvas(buffer, pagesize=letter)
            c.setFont('Helvetica-Bold', 16)
            c.drawString(100, 770, 'REPORTE GERENTE - SIC-NetPOLIx')
            c.setFont('Helvetica-Bold', 12)
            y = 740
            c.drawString(100, y, 'Videos Mejor Calificados')
            c.setFont('Helvetica', 11)
            y -= 20
            for i, r in enumerate(mejor):
                c.drawString(110, y, f'{i+1}. {r[0]}  —  {float(r[1]):.2f}')
                y -= 18
            y -= 10
            c.setFont('Helvetica-Bold', 12)
            c.drawString(100, y, 'Videos Peor Calificados')
            c.setFont('Helvetica', 11)
            y -= 20
            for i, r in enumerate(peor):
                c.drawString(110, y, f'{i+1}. {r[0]}  —  {float(r[1]):.2f}')
                y -= 18
            y -= 10
            c.setFont('Helvetica-Bold', 12)
            c.drawString(100, y, 'Actividad de Clientes (Transacciones)')
            c.setFont('Helvetica', 11)
            y -= 20
            for r in transacciones:
                c.drawString(110, y, f'{r[0]}: {r[1]} transacciones  —  ${float(r[2]):.2f}')
                y -= 18
            y -= 10
            c.setFont('Helvetica-Bold', 12)
            c.drawString(100, y, 'Categorias Mas Populares')
            c.setFont('Helvetica', 11)
            y -= 20
            for r in categorias:
                c.drawString(110, y, f'{r[0]}: {r[1]} transacciones')
                y -= 18
            c.save()
            buffer.seek(0)
            return send_file(buffer, as_attachment=True, download_name='reporte_gerente.pdf',
                             mimetype='application/pdf')

        if formato == 'excel':
            wb = openpyxl.Workbook()

            header_font = Font(bold=True, color='FFFFFF')
            header_fill = PatternFill(fill_type='solid', fgColor='1a1a2e')
            center = Alignment(horizontal='center')

            def crear_hoja(nombre, encabezados, filas):
                ws = wb.create_sheet(nombre)
                for col, titulo in enumerate(encabezados, 1):
                    cell = ws.cell(row=1, column=col, value=titulo)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = center
                for row_idx, fila in enumerate(filas, 2):
                    for col_idx, valor in enumerate(fila, 1):
                        ws.cell(row=row_idx, column=col_idx, value=valor)

            crear_hoja('Mejor Calificados', ['#', 'Titulo', 'Promedio'],
                       [(i+1, r[0], round(float(r[1]), 2)) for i, r in enumerate(mejor)])
            crear_hoja('Peor Calificados', ['#', 'Titulo', 'Promedio'],
                       [(i+1, r[0], round(float(r[1]), 2)) for i, r in enumerate(peor)])
            crear_hoja('Transacciones', ['Tipo', 'Cantidad', 'Total Ingresos ($)'],
                       [(r[0], r[1], round(float(r[2]), 2)) for r in transacciones])
            crear_hoja('Categorias Populares', ['Categoria', 'Transacciones'],
                       [(r[0], r[1]) for r in categorias])

            if 'Sheet' in wb.sheetnames:
                del wb['Sheet']

            buffer = BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            return send_file(buffer, as_attachment=True, download_name='reporte_gerente.xlsx',
                             mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        return {
            'mejor_calificados': [{'titulo': r[0], 'promedio': float(r[1])} for r in mejor],
            'peor_calificados': [{'titulo': r[0], 'promedio': float(r[1])} for r in peor],
            'transacciones': [{'tipo': r[0], 'cantidad': r[1], 'total': float(r[2])} for r in transacciones],
            'categorias_populares': [{'categoria': r[0], 'transacciones': r[1]} for r in categorias]
        }

@api.route('/gerente/descargar/<int:id_video>')
class DescargarVideo(Resource):
    @token_required(roles=['GERENTE'])
    def get(self, id_video):
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT url_descarga, titulo_original FROM video WHERE id_video = ?", (id_video,))
        row = cursor.fetchone()
        conn.close()
        if not row or not row[0]:
            return {'error': 'Video sin URL de descarga'}, 404
        if os.path.exists(row[0]):
            return send_file(row[0], as_attachment=True, download_name=f"{row[1]}.mp4")
        return {'url_descarga': row[0], 'mensaje': 'Puedes descargar desde este enlace'}

# ============================================
# RUTAS DEL FRONTEND
# ============================================

@app.route('/')
def home_redirect():
    return redirect('/login')

@app.route('/login')
def login_page():
    return render_template('index.html')

@app.route('/registro')
def registro_page():
    return render_template('registro.html')

@app.route('/dashboard')
def dashboard_page():
    return render_template('dashboard.html')

@app.route('/gerente')
def gerente_page():
    return render_template('gerente.html')

@app.route('/admin')
def admin_page():
    return render_template('admin.html')


if __name__ == '__main__':
    app.run(debug=True, port=5000)